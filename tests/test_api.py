from __future__ import annotations

import csv
import io

from openpyxl import Workbook


def asset_payload(**overrides):
    payload = {
        "asset_code": "TST-SCADA-001",
        "system_name": "Test supervisory server",
        "site": "Test Site",
        "asset_type": "SCADA Server",
        "support_status": "end_of_support",
        "business_criticality": 5,
        "safety_impact": 4,
        "production_impact": 5,
        "cyber_exposure": 5,
        "failure_likelihood": 4,
        "spares_risk": 4,
        "recoverability_risk": 5,
        "dependency_complexity": 5,
        "evidence_confidence": "A",
        "delivery_readiness": 2,
        "treatment": "contain",
    }
    payload.update(overrides)
    return payload


def live_register_workbook_bytes() -> io.BytesIO:
    workbook = Workbook()
    register = workbook.active
    register.title = "Obsolescence Register"
    headers = [
        "Component Part",
        "Component",
        "Component Description",
        "Manufacturer",
        "Supplier",
        "Responsible Team",
        "Area",
        "Hardware/Software",
        "Functional / Non-Functional",
        "Obsolete (Y/N)",
        "End of Support",
        "Cost Hit (EOS - 1 Year)",
        "Status",
        "Last Obsolescence Check or Update",
        "System Criticality",
        "Obsolescence Criticality",
        "Cost Criticality",
        "Replacement identified (Y/N)",
        "Qty in Field",
        "NR Spares Oty (Battle Boxes)",
        "Telent Spares Qty",
        "Recommend No of Spares",
        "Unit Cost (est)",
        "Total Estimated Cost (Component + Labour)",
        "Full upgrade or ad-hoc replacement?",
        "Risk Score",
        "Risk Factor",
        "Reason for Risk",
        "Risk Assessment",
    ]
    register.append(headers)
    rows = [
        {
            "Component Part": "BIG-AWF-R2800 Reverse Proxy Firewall",
            "Component": "Firewall",
            "Component Description": "Reverse Proxy firewall in the DMZ",
            "Manufacturer": "F5",
            "Supplier": "F5",
            "Responsible Team": "Network",
            "Area": "Network",
            "Hardware/Software": "Software",
            "Functional / Non-Functional": "Functional",
            "Obsolete (Y/N)": "Y",
            "End of Support": "01-Jan-25",
            "Status": "Live / Current",
            "Last Obsolescence Check or Update": "01-Jun-26",
            "System Criticality": 4,
            "Obsolescence Criticality": 3,
            "Cost Criticality": 2,
            "Replacement identified (Y/N)": "Y",
            "Qty in Field": 2,
            "NR Spares Oty (Battle Boxes)": 0,
            "Telent Spares Qty": 0,
            "Recommend No of Spares": 1,
            "Unit Cost (est)": 1000,
            "Total Estimated Cost (Component + Labour)": 2800,
            "Full upgrade or ad-hoc replacement?": "Full Upgrade",
            "Risk Score": 5.0,
            "Risk Factor": "Critical",
            "Reason for Risk": "End of support date has passed",
            "Risk Assessment": "Replace through planned renewal",
        },
        {
            "Component Part": "HIST-SRV-01",
            "Component": "Historian Server",
            "Component Description": "Site historian server",
            "Manufacturer": "Dell",
            "Supplier": "Dell",
            "Responsible Team": "Control",
            "Area": "CMS",
            "Hardware/Software": "Hardware",
            "Functional / Non-Functional": "Non-Functional",
            "Obsolete (Y/N)": "N",
            "Status": "Live / Current",
            "Last Obsolescence Check or Update": "01-Jun-26",
            "System Criticality": 3,
            "Obsolescence Criticality": 2,
            "Cost Criticality": 2,
            "Replacement identified (Y/N)": "N",
            "Qty in Field": 1,
            "NR Spares Oty (Battle Boxes)": 1,
            "Telent Spares Qty": 0,
            "Recommend No of Spares": 1,
            "Unit Cost (est)": 2500,
            "Total Estimated Cost (Component + Labour)": 3500,
            "Full upgrade or ad-hoc replacement?": "Ad-hoc replacement",
            "Risk Score": 3.0,
            "Risk Factor": "Medium",
            "Reason for Risk": "Support date needs confirmation",
            "Risk Assessment": "Monitor and confirm lifecycle evidence",
        },
    ]
    for source_row in rows:
        register.append([source_row.get(header) for header in headers])

    cost_hit_column = headers.index("Cost Hit (EOS - 1 Year)") + 1
    register.cell(row=2, column=cost_hit_column).value = "=K2-365"
    register.cell(row=3, column=cost_hit_column).value = "=K3-365"
    workbook.create_sheet("Front Sheet")
    workbook.create_sheet("Risk Factor Key")
    cross_reference = workbook.create_sheet("PA certs items cross-ref")
    cross_reference.sheet_state = "hidden"

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


def test_health_and_empty_dashboard(client):
    health = client.get("/api/health")
    assert health.status_code == 200
    assert health.json()["database"] == "reachable"

    dashboard = client.get("/api/dashboard/summary").json()
    assert dashboard["generated_from"] == "live_database"
    assert dashboard["metrics"]["total_assets"] == 0


def test_ui_alias_and_favicon_do_not_404(client):
    ui = client.get("/ui")
    assert ui.status_code == 200
    assert "SCADA Lifecycle Command Centre" in ui.text

    favicon = client.get("/favicon.ico")
    assert favicon.status_code == 204


def test_operational_compatibility_endpoints(client):
    health = client.get("/admin/health")
    assert health.status_code == 200
    assert health.headers["Cache-Control"] == "no-store"
    assert health.json()["status"] == "ok"
    assert health.json()["database"] == "reachable"

    trust = client.get("/trust/status")
    assert trust.status_code == 200
    assert trust.headers["Cache-Control"] == "no-store"
    trust_body = trust.json()
    assert trust_body["status"] == "ok"
    assert trust_body["trust_status"] == "trusted"
    assert trust_body["data_source"] == "relational_database"

    metrics = client.get("/admin/metrics")
    assert metrics.status_code == 200
    assert metrics.headers["Cache-Control"] == "no-store"
    assert metrics.headers["content-type"].startswith("text/plain")
    assert "scada_database_reachable 1" in metrics.text
    assert "scada_assets_total 0" in metrics.text
    assert "scada_programmes_total 0" in metrics.text


def test_asset_crud_recalculates_assessment(client):
    created = client.post("/api/assets", json=asset_payload())
    assert created.status_code == 201, created.text
    asset = created.json()
    assert asset["risk_band"] in {"High", "Critical"}
    assert asset["recommended_wave"] == "Wave 0"
    first_score = asset["risk_score"]

    updated = client.patch(
        f"/api/assets/{asset['id']}",
        json={
            "support_status": "supported",
            "cyber_exposure": 1,
            "failure_likelihood": 1,
            "spares_risk": 1,
            "recoverability_risk": 1,
            "dependency_complexity": 1,
            "evidence_confidence": "A",
        },
    )
    assert updated.status_code == 200, updated.text
    changed = updated.json()
    assert changed["risk_score"] < first_score

    listed = client.get("/api/assets", params={"risk_band": changed["risk_band"]}).json()
    assert listed["total"] == 1

    deleted = client.delete(f"/api/assets/{asset['id']}")
    assert deleted.status_code == 200
    assert client.get("/api/assets").json()["total"] == 0


def test_asset_validation_rejects_blank_identity_and_invalid_programme(client):
    created = client.post("/api/assets", json=asset_payload())
    assert created.status_code == 201, created.text
    asset = created.json()

    blank_update = client.patch(f"/api/assets/{asset['id']}", json={"asset_code": "   "})
    assert blank_update.status_code == 422

    null_update = client.patch(f"/api/assets/{asset['id']}", json={"site": None})
    assert null_update.status_code == 422

    unchanged = client.get(f"/api/assets/{asset['id']}")
    assert unchanged.status_code == 200
    assert unchanged.json()["asset_code"] == "TST-SCADA-001"

    zero_programme = client.post(
        "/api/assets",
        json=asset_payload(asset_code="TST-SCADA-002", programme_id=0),
    )
    assert zero_programme.status_code == 422

    missing_programme = client.post(
        "/api/assets",
        json=asset_payload(asset_code="TST-SCADA-003", programme_id=999),
    )
    assert missing_programme.status_code == 400
    assert missing_programme.json()["detail"] == "programme_id does not exist"


def test_programme_linkage_and_dashboard_rollup(client):
    programme_response = client.post(
        "/api/programmes",
        json={
            "package_code": "OBP-T01",
            "title": "Test renewal package",
            "status": "design",
            "target_wave": "Wave 1",
            "budget_estimate": 100000,
            "contingency_percent": 25,
        },
    )
    assert programme_response.status_code == 201, programme_response.text
    programme = programme_response.json()

    asset = client.post(
        "/api/assets",
        json=asset_payload(programme_id=programme["id"]),
    ).json()
    package = client.get(f"/api/programmes/{programme['id']}").json()
    assert package["asset_count"] == 1
    assert package["total_budget"] == 125000
    assert package["average_risk"] == asset["risk_score"]

    dashboard = client.get("/api/dashboard/summary").json()
    assert dashboard["metrics"]["active_programmes"] == 1
    assert dashboard["metrics"]["funding_estimate"] == 125000


def test_programme_validation_rejects_blank_identity(client):
    blank = client.post("/api/programmes", json={"package_code": "   ", "title": "   "})
    assert blank.status_code == 422

    created = client.post(
        "/api/programmes",
        json={"package_code": "OBP-T03", "title": "Validation package"},
    )
    assert created.status_code == 201, created.text

    blank_patch = client.patch(f"/api/programmes/{created.json()['id']}", json={"title": "   "})
    assert blank_patch.status_code == 422


def test_csv_import_uses_row_savepoints_and_export(client):
    programme = client.post(
        "/api/programmes",
        json={"package_code": "OBP-T02", "title": "Imported package"},
    ).json()
    assert programme["package_code"] == "OBP-T02"

    csv_text = """asset_code,system_name,site,business_criticality,support_status,programme_code
IMP-001,Imported PLC,North,4,limited,OBP-T02
IMP-BAD,Bad rating,North,9,unknown,OBP-T02
IMP-002,Imported RTU,South,3,supported,OBP-T02
"""
    response = client.post(
        "/api/assets/import",
        files={"file": ("register.csv", io.BytesIO(csv_text.encode()), "text/csv")},
    )
    assert response.status_code == 200, response.text
    result = response.json()
    assert result["created"] == 2
    assert result["failed"] == 1
    assert client.get("/api/assets").json()["total"] == 2

    exported = client.get("/api/assets/export.csv")
    assert exported.status_code == 200
    assert "IMP-001" in exported.text
    assert "IMP-002" in exported.text
    assert "IMP-BAD" not in exported.text
    assert "OBP-T02" in exported.text


def test_live_obsolescence_register_import_mapping_and_upsert(client):
    csv_text = """Component Part,Component,Component Description,Manufacturer,Responsible Team,Area,Hardware/Software,Functional / Non-Functional,Obsolete (Y/N),End of Support,Status,System Criticality,Obsolescence Criticality,Cost Criticality,Replacement identified (Y/N),Full upgrade or ad-hoc replacement?,Risk Score,Risk Factor,Reason for Risk
BIG-AWF-R2800 Reverse Proxy Firewall,Firewall,Reverse Proxy firewall in the DMZ,F5,Network,Network,Software,Functional,Y,01-Jan-25,Live / Current,4,3,2,Y,Full Upgrade,5.00,Critical,End of support date has passed
BIG-AWF-R2800 Reverse Proxy Firewall,Firewall,Reverse Proxy firewall in the DMZ,F5,Network,Network,Software,Functional,Y,01-Jan-25,Live / Current,4,3,2,Y,Full Upgrade,5.00,Critical,End of support date has passed
"""
    response = client.post(
        "/api/assets/import",
        files={"file": ("live-register.csv", io.BytesIO(csv_text.encode()), "text/csv")},
    )
    assert response.status_code == 200, response.text
    result = response.json()
    assert result["created"] == 2
    assert result["updated"] == 0
    assert result["failed"] == 0

    listed = client.get("/api/assets", params={"limit": 10, "sort_by": "asset_code", "sort_dir": "asc"}).json()
    assert listed["total"] == 2
    asset_codes = {asset["asset_code"] for asset in listed["items"]}
    assert len(asset_codes) == 2
    for asset in listed["items"]:
        assert asset["asset_code"].startswith("TPCMS-BIG-AWF-R2800-REVERSE-PROXY-FIREWALL-")
        assert asset["system_name"] == "Firewall"
        assert asset["site"] == "Network"
        assert asset["process_area"] == "Network"
        assert asset["asset_type"] == "Firewall"
        assert asset["manufacturer"] == "F5"
        assert asset["model"] == "BIG-AWF-R2800 Reverse Proxy Firewall"
        assert asset["support_status"] == "end_of_support"
        assert asset["treatment"] == "replace"
        assert "Original risk factor: Critical" in asset["notes"]

    repeated = client.post(
        "/api/assets/import",
        files={"file": ("live-register.csv", io.BytesIO(csv_text.encode()), "text/csv")},
    )
    assert repeated.status_code == 200, repeated.text
    repeated_result = repeated.json()
    assert repeated_result["created"] == 0
    assert repeated_result["updated"] == 2
    assert repeated_result["failed"] == 0


def test_live_obsolescence_workbook_import_preserves_source_fields_and_exports(client):
    response = client.post(
        "/api/assets/import",
        files={
            "file": (
                "live-register.xlsx",
                live_register_workbook_bytes(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert response.status_code == 200, response.text
    result = response.json()
    assert result["created"] == 2
    assert result["updated"] == 0
    assert result["failed"] == 0

    listed = client.get("/api/assets", params={"source_category": "Network", "limit": 10}).json()
    assert listed["total"] == 1
    asset = listed["items"][0]
    assert asset["source_workbook"] == "live-register.xlsx"
    assert asset["source_sheet"] == "Obsolescence Register"
    assert asset["source_row"] == 2
    assert asset["source_category"] == "Network"
    assert asset["source_subcategory"] == "Network"
    assert asset["source_payload"]["Component Part"] == "BIG-AWF-R2800 Reverse Proxy Firewall"
    assert asset["source_payload"]["Total Estimated Cost (Component + Labour)"] == 2800
    assert asset["source_formulas"]["Cost Hit (EOS - 1 Year)"] == "=K2-365"
    assert asset["source_intelligence"]["source_risk_factor"] == "Critical"
    assert asset["source_intelligence"]["collation"]["hardware_software"] == "Software"
    assert "source_high_risk" in asset["source_intelligence"]["automation_flags"]

    source_summary = client.get("/api/dashboard/source-summary")
    assert source_summary.status_code == 200
    source = source_summary.json()
    assert source["source_records"] == 2
    assert source["source_workbooks"] == ["live-register.xlsx"]
    assert source["totals"]["estimated_cost"] == 6300
    assert source["totals"]["quantity_in_field"] == 3
    assert source["totals"]["formula_columns"] == 1
    assert {"name": "Network", "count": 1} in source["source_categories"]
    assert {"name": "Cost Hit (EOS - 1 Year)", "count": 2} in source["formula_columns"]

    source_export = client.get("/api/assets/source-export.csv")
    assert source_export.status_code == 200
    exported_rows = list(csv.DictReader(io.StringIO(source_export.text)))
    assert len(exported_rows) == 2
    network_row = next(row for row in exported_rows if row["Area"] == "Network")
    assert network_row["Component Part"] == "BIG-AWF-R2800 Reverse Proxy Firewall"
    assert network_row["platform_source_category"] == "Network"
    assert network_row["platform_risk_band"] in {"High", "Critical"}
    assert network_row["platform_formula_columns"] == "Cost Hit (EOS - 1 Year)"
    assert "source_high_risk" in network_row["platform_automation_flags"]

    repeated = client.post(
        "/api/assets/import",
        files={
            "file": (
                "live-register.xlsx",
                live_register_workbook_bytes(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert repeated.status_code == 200, repeated.text
    repeated_result = repeated.json()
    assert repeated_result["created"] == 0
    assert repeated_result["updated"] == 2
    assert repeated_result["failed"] == 0


def test_csv_import_rejects_unknown_direct_programme_id(client):
    csv_text = """asset_code,system_name,site,programme_id
IMP-NO-PROG,Imported orphan,North,999
"""
    response = client.post(
        "/api/assets/import",
        files={"file": ("register.csv", io.BytesIO(csv_text.encode()), "text/csv")},
    )
    assert response.status_code == 200, response.text
    result = response.json()
    assert result["created"] == 0
    assert result["failed"] == 1
    assert "programme_id does not exist" in result["errors"][0]
    assert client.get("/api/assets").json()["total"] == 0
