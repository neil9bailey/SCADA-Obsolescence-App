from __future__ import annotations

from datetime import date, datetime
from hashlib import sha1
from io import BytesIO
from pathlib import Path
import re
from typing import Any

from openpyxl import load_workbook

NULL_MARKERS = {"", "-", "N/A", "#N/A", "TBC", "NOT ANNOUNCED", "NO DATE GIVEN"}
LIVE_REGISTER_SHEET = "Obsolescence Register"

LIVE_REGISTER_REQUIRED_COLUMNS = {
    "Component Part",
    "Component",
    "Area",
    "End of Support",
    "Risk Factor",
}

LIVE_REGISTER_IDENTITY_COLUMNS = (
    "Component Part",
    "Component",
    "Area",
    "Manufacturer",
    "Hardware/Software",
    "Component Description",
)


def _clean(value: Any) -> str | None:
    if value is None:
        return None
    text = " ".join(str(value).strip().split())
    return None if text.upper() in NULL_MARKERS else text


def _json_value(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if value is None:
        return None
    if isinstance(value, str):
        return value
    if isinstance(value, (int, float, bool)):
        return value
    return str(value)


def _truncate(value: str | None, max_length: int) -> str | None:
    if value is None:
        return None
    return value[:max_length]


def _parse_number(value: Any) -> float | None:
    if isinstance(value, (int, float)):
        return float(value)
    text = _clean(value)
    if text is None:
        return None
    match = re.search(r"-?\d+(?:\.\d+)?", text.replace(",", ""))
    return float(match.group(0)) if match else None


def _parse_int(value: Any) -> int | None:
    number = _parse_number(value)
    return int(round(number)) if number is not None else None


def _rating(value: Any, default: int) -> int:
    parsed = _parse_number(value)
    if parsed is None:
        return default
    return max(1, min(5, int(round(parsed))))


def _source_risk_rating(value: Any, default: int = 3) -> int:
    text = (_clean(value) or "").lower()
    if "critical" in text:
        return 5
    if "high" in text:
        return 4
    if "medium" in text:
        return 3
    if "low" in text:
        return 2
    if "no component" in text:
        return 1
    return default


def _parse_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _clean(value)
    if text is None:
        return None
    for fmt in ("%d-%b-%y", "%d-%b-%Y", "%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _slug(value: str | None) -> str:
    text = value or "row"
    slug = re.sub(r"[^A-Za-z0-9]+", "-", text.upper()).strip("-")
    return (slug or "ROW")[:44]


def normalise_live_header(header: str | None) -> str:
    return (header or "").strip().lstrip("\ufeff")


def is_live_register_fieldnames(fieldnames: list[str] | None) -> bool:
    cleaned = {normalise_live_header(name) for name in fieldnames or []}
    return LIVE_REGISTER_REQUIRED_COLUMNS.issubset(cleaned)


def live_register_natural_key(row: dict[str, Any]) -> tuple[str, ...]:
    return tuple(_clean(row.get(column)) or "" for column in LIVE_REGISTER_IDENTITY_COLUMNS)


def _normalised_source_payload(row: dict[str, Any]) -> dict[str, Any]:
    return {
        normalise_live_header(key): _json_value(value)
        for key, value in row.items()
        if normalise_live_header(key)
    }


def _asset_code(row: dict[str, Any], duplicate_index: int) -> str:
    part = _clean(row.get("Component Part"))
    component = _clean(row.get("Component"))
    natural_key = live_register_natural_key(row)
    hash_basis = "|".join(natural_key)
    if duplicate_index > 1:
        hash_basis = f"{hash_basis}|duplicate:{duplicate_index}"
    digest = sha1(hash_basis.encode("utf-8")).hexdigest()[:8].upper()
    return f"TPCMS-{_slug(part or component)}-{digest}"[:80]


def _support_status(row: dict[str, Any], support_end_date: date | None, today: date | None = None) -> str:
    today = today or date.today()
    status = (_clean(row.get("Status")) or "").lower()
    risk_factor = (_clean(row.get("Risk Factor")) or "").lower()
    obsolete = (_clean(row.get("Obsolete (Y/N)")) or "").upper()
    raw_eos_text = str(row.get("End of Support") or "").strip().upper()

    if "decom" in status or "no component" in risk_factor:
        return "supported"
    if support_end_date is not None:
        return "end_of_support" if support_end_date <= today else "limited"
    if obsolete == "Y":
        return "end_of_support"
    known_no_eos = {"N/A", "#N/A", "-", "NOT ANNOUNCED", "NO DATE GIVEN"}
    if obsolete == "N" and raw_eos_text in known_no_eos:
        return "supported"
    return "unknown"


def _spares_risk(row: dict[str, Any], fallback: int) -> int:
    recommended = _parse_int(row.get("Recommend No of Spares"))
    nr_spares = _parse_int(row.get("NR Spares Oty (Battle Boxes)")) or 0
    telent_spares = _parse_int(row.get("Telent Spares Qty")) or 0
    in_field = _parse_int(row.get("Qty in Field")) or 0
    available = nr_spares + telent_spares

    if recommended is not None:
        if available < recommended:
            return 5
        if available == recommended:
            return 3
        return 2
    if in_field > 0 and available == 0:
        return 4
    return fallback


def _cyber_exposure(row: dict[str, Any], fallback: int) -> int:
    area = (_clean(row.get("Area")) or "").lower()
    hardware_software = (_clean(row.get("Hardware/Software")) or "").lower()
    component = (_clean(row.get("Component")) or "").lower()
    if hardware_software == "software" or area in {"network", "security", "cms", "voice"}:
        return max(4, fallback)
    if any(keyword in component for keyword in ("firewall", "router", "switch", "server")):
        return max(4, fallback)
    return fallback


def _evidence_confidence(row: dict[str, Any]) -> str:
    has_recent_check = _parse_date(row.get("Last Obsolescence Check or Update")) is not None
    numeric_criticality = all(
        _parse_number(row.get(column)) is not None
        for column in ("System Criticality", "Obsolescence Criticality", "Cost Criticality")
    )
    if has_recent_check and numeric_criticality:
        return "A"
    if has_recent_check or numeric_criticality:
        return "B"
    if _clean(row.get("Risk Factor")):
        return "C"
    return "D"


def _delivery_readiness(row: dict[str, Any]) -> int:
    status = (_clean(row.get("Status")) or "").lower()
    replacement = (_clean(row.get("Replacement identified (Y/N)")) or "").upper()
    if "upgrade instructed" in status or "in progress" in status:
        return 4
    if "obsolescence" in status:
        return 3
    if replacement == "Y":
        return 2
    return 1


def _treatment(row: dict[str, Any]) -> str:
    status = (_clean(row.get("Status")) or "").lower()
    upgrade_type = (_clean(row.get("Full upgrade or ad-hoc replacement?")) or "").lower()
    obsolete = (_clean(row.get("Obsolete (Y/N)")) or "").upper()
    replacement = (_clean(row.get("Replacement identified (Y/N)")) or "").upper()
    risk_factor = (_clean(row.get("Risk Factor")) or "").lower()

    if "decom" in status or "no component" in risk_factor:
        return "decommission"
    if "full upgrade" in upgrade_type:
        return "replace"
    if "ad-hoc" in upgrade_type or "ad hoc" in upgrade_type:
        return "upgrade"
    if obsolete == "Y" and replacement == "Y":
        return "replace"
    if obsolete == "Y":
        return "contain"
    if replacement == "Y":
        return "upgrade"
    return "assess"


def _notes(row: dict[str, Any], line_number: int) -> str:
    note_fields = (
        ("Source register row", str(line_number)),
        ("Component part", row.get("Component Part")),
        ("Component description", row.get("Component Description")),
        ("Supplier", row.get("Supplier")),
        ("Last obsolescence check/update", row.get("Last Obsolescence Check or Update")),
        ("Obsolete flag", row.get("Obsolete (Y/N)")),
        ("Obsolete date", row.get("Obsolete Date")),
        ("Last buy", row.get("Last Buy")),
        ("End of support", row.get("End of Support")),
        ("Cost hit date", row.get("Cost Hit (EOS - 1 Year)")),
        ("Source status", row.get("Status")),
        ("Obsolescence management strategy", row.get("Obsolescence Management Strategy")),
        ("Quantity in field", row.get("Qty in Field")),
        ("NR spares quantity", row.get("NR Spares Oty (Battle Boxes)")),
        ("Telent spares quantity", row.get("Telent Spares Qty")),
        ("Recommended spares", row.get("Recommend No of Spares")),
        ("Future installation quantity", row.get("Future Installation Qty")),
        ("Spares strategy", row.get("Spares Strategy")),
        ("Unit cost estimate", row.get("Unit Cost (est)")),
        ("Total devices cost", row.get("Cost (of Total Devices in Field)")),
        ("Procurement lead time days", row.get("Estimated Procurement Lead Time (Days)")),
        ("Estimated labour cost", row.get("Estimated Labour Cost")),
        ("Total estimated cost", row.get("Total Estimated Cost (Component + Labour)")),
        ("PA subsystem certificate", row.get("PA Subsystem Cert Reference")),
        ("PA product certificate", row.get("PA (Product Specific) Certificate Reference")),
        ("Upgrade or replacement type", row.get("Full upgrade or ad-hoc replacement?")),
        ("Source system criticality", row.get("System Criticality")),
        ("Source obsolescence criticality", row.get("Obsolescence Criticality")),
        ("Source cost criticality", row.get("Cost Criticality")),
        ("Source criticality rating", row.get("Ctriticality Rating")),
        ("Replacement identified", row.get("Replacement identified (Y/N)")),
        ("Replacement", row.get("Replacement")),
        ("Supporting notes", row.get("Supporting Notes")),
        ("Original risk score", row.get("Risk Score")),
        ("Original risk factor", row.get("Risk Factor")),
        ("Original risk reason", row.get("Reason for Risk")),
        ("Original risk assessment", row.get("Risk Assessment")),
    )
    return "\n".join(
        f"{label}: {cleaned}"
        for label, value in note_fields
        if (cleaned := _clean(value)) is not None
    )


def _source_intelligence(row: dict[str, Any], support_end_date: date | None, spares_risk: int) -> dict[str, Any]:
    source_risk_score = _parse_number(row.get("Risk Score"))
    source_risk_factor = _clean(row.get("Risk Factor"))
    replacement = _clean(row.get("Replacement identified (Y/N)"))
    support_status = _support_status(row, support_end_date)
    support_evidence = "dated" if support_end_date else "missing"
    if _clean(row.get("End of Support")) is not None and support_end_date is None:
        support_evidence = "textual"

    flags: list[str] = []
    if source_risk_factor in {"Critical", "High"}:
        flags.append("source_high_risk")
    if support_status == "end_of_support" and (replacement or "").upper() != "Y":
        flags.append("end_of_support_without_confirmed_replacement")
    if spares_risk >= 5:
        flags.append("spares_below_recommended")
    if not _clean(row.get("Responsible Team")):
        flags.append("missing_owner")
    if support_end_date is None:
        flags.append("missing_support_date")
    if "no component" in (_clean(row.get("Risk Factor")) or "").lower():
        flags.append("no_component_in_field")
    if source_risk_score is None:
        flags.append("missing_source_risk_score")

    return {
        "source_risk_score": source_risk_score,
        "source_risk_factor": source_risk_factor,
        "source_reason_for_risk": _clean(row.get("Reason for Risk")),
        "source_risk_assessment": _clean(row.get("Risk Assessment")),
        "source_status": _clean(row.get("Status")),
        "source_support_evidence": support_evidence,
        "source_replacement_identified": replacement,
        "source_total_estimated_cost": _parse_number(row.get("Total Estimated Cost (Component + Labour)")),
        "source_quantity_in_field": _parse_int(row.get("Qty in Field")),
        "source_unit_cost": _parse_number(row.get("Unit Cost (est)")),
        "collation": {
            "area": _clean(row.get("Area")),
            "responsible_team": _clean(row.get("Responsible Team")),
            "hardware_software": _clean(row.get("Hardware/Software")),
            "functional_class": _clean(row.get("Functional / Non-Functional")),
            "status": _clean(row.get("Status")),
            "risk_factor": source_risk_factor,
            "replacement_identified": replacement,
            "manufacturer": _clean(row.get("Manufacturer")),
            "supplier": _clean(row.get("Supplier")),
        },
        "automation_flags": flags,
    }


def read_live_register_workbook(contents: bytes, filename: str) -> list[dict[str, Any]]:
    values_workbook = load_workbook(BytesIO(contents), data_only=True, read_only=False)
    formula_workbook = load_workbook(BytesIO(contents), data_only=False, read_only=False)
    if LIVE_REGISTER_SHEET not in values_workbook.sheetnames:
        raise ValueError(f"workbook must contain a '{LIVE_REGISTER_SHEET}' sheet")

    values_sheet = values_workbook[LIVE_REGISTER_SHEET]
    formula_sheet = formula_workbook[LIVE_REGISTER_SHEET]
    headers: list[tuple[int, str]] = []
    for column in range(1, values_sheet.max_column + 1):
        header = normalise_live_header(values_sheet.cell(1, column).value)
        if header:
            headers.append((column, header))

    if not is_live_register_fieldnames([header for _, header in headers]):
        raise ValueError("workbook register sheet is missing required live obsolescence columns")

    rows: list[dict[str, Any]] = []
    for row_number in range(2, values_sheet.max_row + 1):
        row: dict[str, Any] = {}
        source_payload: dict[str, Any] = {}
        source_formulas: dict[str, str] = {}
        has_data = False
        for column, header in headers:
            value = values_sheet.cell(row_number, column).value
            formula = formula_sheet.cell(row_number, column).value
            row[header] = value
            source_payload[header] = _json_value(value)
            if value not in (None, ""):
                has_data = True
            if isinstance(formula, str) and formula.startswith("="):
                source_formulas[header] = formula
        if has_data:
            rows.append(
                {
                    "row": row,
                    "line_number": row_number,
                    "source_workbook": Path(filename).name,
                    "source_sheet": LIVE_REGISTER_SHEET,
                    "source_row": row_number,
                    "source_payload": source_payload,
                    "source_formulas": source_formulas,
                }
            )
    return rows


def map_live_register_row(
    row: dict[str, Any],
    *,
    line_number: int,
    duplicate_index: int = 1,
    source_workbook: str | None = None,
    source_sheet: str | None = None,
    source_row: int | None = None,
    source_payload: dict[str, Any] | None = None,
    source_formulas: dict[str, Any] | None = None,
    today: date | None = None,
) -> dict[str, Any]:
    source_risk = _source_risk_rating(row.get("Risk Factor"))
    system_criticality = _rating(row.get("System Criticality"), source_risk)
    obsolescence_criticality = _rating(row.get("Obsolescence Criticality"), source_risk)
    cost_criticality = _rating(row.get("Cost Criticality"), 3)
    support_end_date = _parse_date(row.get("End of Support"))
    functional_type = (_clean(row.get("Functional / Non-Functional")) or "").lower()
    safety_impact = (
        system_criticality if functional_type == "functional" else max(1, system_criticality - 1)
    )
    component = _clean(row.get("Component"))
    description = _clean(row.get("Component Description"))
    hardware_software = _clean(row.get("Hardware/Software"))

    spares_risk = _spares_risk(row, obsolescence_criticality)
    source_intelligence = _source_intelligence(row, support_end_date, spares_risk)
    source_payload = source_payload or _normalised_source_payload(row)
    source_formulas = source_formulas or {}

    return {
        "asset_code": _asset_code(row, duplicate_index),
        "system_name": _truncate(
            component or description or _clean(row.get("Component Part")) or "Unknown component",
            200,
        ),
        "site": _truncate(_clean(row.get("Area")) or "TPCMS", 120),
        "process_area": _truncate(_clean(row.get("Responsible Team")) or _clean(row.get("Area")), 160),
        "asset_type": _truncate(component or hardware_software or "Other", 80),
        "manufacturer": _truncate(_clean(row.get("Manufacturer")) or _clean(row.get("Supplier")), 120),
        "model": _truncate(_clean(row.get("Component Part")), 120),
        "support_status": _support_status(row, support_end_date, today=today),
        "support_end_date": support_end_date,
        "business_criticality": system_criticality,
        "safety_impact": safety_impact,
        "production_impact": system_criticality,
        "cyber_exposure": _cyber_exposure(row, obsolescence_criticality),
        "failure_likelihood": obsolescence_criticality,
        "spares_risk": spares_risk,
        "recoverability_risk": max(obsolescence_criticality, cost_criticality),
        "dependency_complexity": max(2, round((system_criticality + obsolescence_criticality) / 2)),
        "evidence_confidence": _evidence_confidence(row),
        "delivery_readiness": _delivery_readiness(row),
        "treatment": _treatment(row),
        "owner": _truncate(_clean(row.get("Responsible Team")), 120),
        "notes": _notes(row, line_number),
        "source_workbook": _truncate(source_workbook, 260),
        "source_sheet": _truncate(source_sheet, 120),
        "source_row": source_row or line_number,
        "source_category": _truncate(source_intelligence["collation"].get("area"), 80),
        "source_subcategory": _truncate(source_intelligence["collation"].get("responsible_team"), 120),
        "source_payload": source_payload,
        "source_formulas": source_formulas,
        "source_intelligence": source_intelligence,
    }
